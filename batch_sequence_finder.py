#!/usr/bin/env python3
"""
Batch DNA/RNA Sequence Pattern Finder

Description:
    Runs multiple pattern search tasks with different parameters and merges results.
    Accepts JSON configuration file for batch processing with flexible output modes.

Author: Qiande, Moni
Email: mqiande@ufl.edu
License: Apache-2.0
Version: 1.0.0
Created: 2025
Last Modified: 2025-10-12

Usage:
    python3 batch_sequence_finder.py config.json

For more information, see README.md
"""

import argparse
import json
import sys
from typing import List, Dict
from Bio import SeqIO
from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.cell.text import InlineFont
from openpyxl.cell.rich_text import TextBlock, CellRichText


def parse_fasta(fasta_file: str) -> List[tuple]:
    """Parse FASTA file and return list of (id, sequence) tuples."""
    sequences = []
    try:
        for record in SeqIO.parse(fasta_file, "fasta"):
            sequences.append((record.id, str(record.seq).upper()))
        print(f"Loaded {len(sequences)} sequences from {fasta_file}")
        return sequences
    except Exception as e:
        print(f"Error reading FASTA file: {e}")
        sys.exit(1)


def count_mismatches(sequence: str, pattern: str) -> tuple:
    """Count mismatches between sequence and pattern, treating 'n' as wildcard."""
    if len(sequence) != len(pattern):
        return (float('inf'), [])

    mismatches = 0
    mismatch_positions = []

    for i, (seq_char, pat_char) in enumerate(zip(sequence, pattern)):
        if pat_char.lower() == 'n':
            continue
        elif seq_char != pat_char:
            mismatches += 1
            mismatch_positions.append(i)

    return mismatches, mismatch_positions


def build_pattern_with_gaps(patterns: List[str], gap_size: int) -> str:
    """Build a combined pattern with specified gap size between sub-patterns."""
    gap = 'n' * gap_size
    return gap.join(patterns)


def search_pattern_in_sequence(
    seq_id: str,
    sequence: str,
    patterns: List[str],
    gap_size: int,
    max_mismatches: int,
    task_name: str
) -> List[Dict]:
    """Search for pattern matches in a sequence with specified gap size."""
    results = []

    full_pattern = build_pattern_with_gaps(patterns, gap_size)
    pattern_length = len(full_pattern)
    seq_length = len(sequence)

    for i in range(len(sequence) - pattern_length + 1):
        subseq = sequence[i:i + pattern_length]
        mismatch_count, mismatch_positions = count_mismatches(subseq, full_pattern)

        if mismatch_count <= max_mismatches:
            start_pos = i + 1  # 1-based indexing
            tss_position = start_pos - seq_length  # Position relative to sequence end

            results.append({
                'task_name': task_name,
                'seq_id': seq_id,
                'subsequence': subseq,
                'start_pos': start_pos,
                'tss_position': tss_position,
                'mismatch_count': mismatch_count,
                'gap_size': gap_size,
                'pattern': ', '.join(patterns),
                'mismatch_positions': mismatch_positions
            })

    return results


def create_rich_text_cell(text: str, mismatch_positions: List[int]) -> CellRichText:
    """Create rich text cell with red-colored characters at mismatch positions."""
    if not mismatch_positions:
        return text

    text_blocks = []
    mismatch_set = set(mismatch_positions)

    i = 0
    while i < len(text):
        if i in mismatch_set:
            j = i
            while j < len(text) and j in mismatch_set:
                j += 1
            text_blocks.append(
                TextBlock(InlineFont(color='FF0000'), text[i:j])
            )
            i = j
        else:
            j = i
            while j < len(text) and j not in mismatch_set:
                j += 1
            text_blocks.append(
                TextBlock(InlineFont(color='000000'), text[i:j])
            )
            i = j

    return CellRichText(*text_blocks)


def write_results_to_excel(results: List[Dict], output_file: str, mode: str = 'single'):
    """
    Write results to Excel file with colored mismatches.

    Args:
        results: List of match dictionaries
        output_file: Path to output Excel file
        mode: 'single' for one sheet, 'multi' for separate sheets per task
    """
    wb = Workbook()
    wb.remove(wb.active)  # Remove default sheet

    if mode == 'single':
        # Single sheet with all results
        ws = wb.create_sheet("All Results")

        headers = ['Task Name', 'Sequence ID', 'Found Subsequence', 'Start Position',
                   'Transcript Start Site Position', 'Mismatch Count', 'Gap Size', 'Pattern']
        ws.append(headers)

        # Style header
        header_font = Font(bold=True)
        for cell in ws[1]:
            cell.font = header_font

        # Write all results
        for result in results:
            row = [
                result['task_name'],
                result['seq_id'],
                create_rich_text_cell(result['subsequence'], result['mismatch_positions']),
                result['start_pos'],
                result['tss_position'],
                result['mismatch_count'],
                result['gap_size'],
                result['pattern']
            ]
            ws.append(row)

        # Auto-adjust column widths
        for column in ws.columns:
            max_length = 0
            column_letter = column[0].column_letter
            for cell in column:
                try:
                    if cell.value:
                        max_length = max(max_length, len(str(cell.value)))
                except:
                    pass
            adjusted_width = min(max_length + 2, 50)
            ws.column_dimensions[column_letter].width = adjusted_width

    else:  # multi-sheet mode
        # Group results by task
        task_results = {}
        for result in results:
            task_name = result['task_name']
            if task_name not in task_results:
                task_results[task_name] = []
            task_results[task_name].append(result)

        # Create sheet for each task
        for task_name, task_data in task_results.items():
            # Sanitize sheet name (Excel has restrictions)
            sheet_name = task_name[:31]  # Max 31 chars
            ws = wb.create_sheet(sheet_name)

            headers = ['Sequence ID', 'Found Subsequence', 'Start Position',
                       'Transcript Start Site Position', 'Mismatch Count', 'Gap Size', 'Pattern']
            ws.append(headers)

            # Style header
            header_font = Font(bold=True)
            for cell in ws[1]:
                cell.font = header_font

            # Write results for this task
            for result in task_data:
                row = [
                    result['seq_id'],
                    create_rich_text_cell(result['subsequence'], result['mismatch_positions']),
                    result['start_pos'],
                    result['tss_position'],
                    result['mismatch_count'],
                    result['gap_size'],
                    result['pattern']
                ]
                ws.append(row)

            # Auto-adjust column widths
            for column in ws.columns:
                max_length = 0
                column_letter = column[0].column_letter
                for cell in column:
                    try:
                        if cell.value:
                            max_length = max(max_length, len(str(cell.value)))
                    except:
                        pass
                adjusted_width = min(max_length + 2, 50)
                ws.column_dimensions[column_letter].width = adjusted_width

    wb.save(output_file)
    print(f"\nResults written to {output_file}")


def load_config(config_file: str) -> Dict:
    """Load and validate JSON configuration file."""
    try:
        with open(config_file, 'r') as f:
            config = json.load(f)

        # Validate required fields
        if 'tasks' not in config:
            print("Error: Config must contain 'tasks' field")
            sys.exit(1)

        if 'input_fasta' not in config:
            print("Error: Config must contain 'input_fasta' field")
            sys.exit(1)

        if 'output_excel' not in config:
            print("Error: Config must contain 'output_excel' field")
            sys.exit(1)

        # Set defaults
        if 'output_mode' not in config:
            config['output_mode'] = 'single'

        # Validate each task
        for i, task in enumerate(config['tasks']):
            if 'patterns' not in task:
                print(f"Error: Task {i} must contain 'patterns' field")
                sys.exit(1)

            # Set task defaults
            if 'name' not in task:
                task['name'] = f"Task_{i+1}"
            if 'gap_sizes' not in task:
                task['gap_sizes'] = [0, 1, 2, 3]
            if 'max_mismatches' not in task:
                task['max_mismatches'] = 2

        return config

    except FileNotFoundError:
        print(f"Error: Config file '{config_file}' not found")
        sys.exit(1)
    except json.JSONDecodeError as e:
        print(f"Error: Invalid JSON in config file: {e}")
        sys.exit(1)


def main():
    """Main function to run batch sequence pattern finder."""
    parser = argparse.ArgumentParser(
        description='Batch search for DNA/RNA sequence patterns using JSON config',
        formatter_class=argparse.RawDescriptionHelpFormatter,
        epilog="""
Example config.json:
{
  "input_fasta": "input.fasta",
  "output_excel": "results.xlsx",
  "output_mode": "single",
  "tasks": [
    {
      "name": "Task1_GAA_TTC",
      "patterns": ["nGAAn", "nTTCn", "nGAAn"],
      "gap_sizes": [0, 1, 2, 3],
      "max_mismatches": 2
    },
    {
      "name": "Task2_Custom",
      "patterns": ["ACGT", "TGCA"],
      "gap_sizes": [0, 1],
      "max_mismatches": 1
    }
  ]
}

Usage:
  python batch_sequence_finder.py config.json
        """
    )

    parser.add_argument('config', help='JSON configuration file')

    args = parser.parse_args()

    # Load configuration
    config = load_config(args.config)

    print(f"Loading configuration from {args.config}")
    print(f"Input FASTA: {config['input_fasta']}")
    print(f"Output Excel: {config['output_excel']}")
    print(f"Output mode: {config['output_mode']}")
    print(f"Number of tasks: {len(config['tasks'])}\n")

    # Parse FASTA file
    sequences = parse_fasta(config['input_fasta'])

    # Run all tasks
    all_results = []

    for task in config['tasks']:
        task_name = task['name']
        patterns = [p.upper() for p in task['patterns']]
        gap_sizes = task['gap_sizes']
        max_mismatches = task['max_mismatches']

        print(f"Running task: {task_name}")
        print(f"  Patterns: {patterns}")
        print(f"  Gap sizes: {gap_sizes}")
        print(f"  Max mismatches: {max_mismatches}")

        task_results = []

        for gap_size in gap_sizes:
            for seq_id, sequence in sequences:
                results = search_pattern_in_sequence(
                    seq_id, sequence, patterns, gap_size, max_mismatches, task_name
                )
                task_results.extend(results)

        print(f"  Found {len(task_results)} matches\n")
        all_results.extend(task_results)

    print(f"Total matches across all tasks: {len(all_results)}")

    # Write results to Excel
    if all_results:
        write_results_to_excel(all_results, config['output_excel'], config['output_mode'])
    else:
        print("No matches found. No output file created.")


if __name__ == '__main__':
    main()
